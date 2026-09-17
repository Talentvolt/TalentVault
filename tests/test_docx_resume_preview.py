"""
Live DOCX resume preview tests.

Verifies:
- DOCX resumes render server-side (Mammoth, pure Python) into the Resume
  Preview area instead of the "Preview not available for this file type (DOCX)"
  message.
- Basic formatting is preserved.
- A genuinely corrupted DOCX falls back gracefully (HTTP 200, not a 500).
- PDF preview is unchanged.
"""
import io
import pathlib

import docx
import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse

from apps.candidates.models import CandidateProfile

User = get_user_model()

DOCX_CT = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
PDF_CT = 'application/pdf'


def _docx_bytes():
    buf = io.BytesIO()
    d = docx.Document()
    d.add_heading('Jane Smith', level=1)
    p = d.add_paragraph()
    run = p.add_run('Senior Engineer with Python and AWS experience.')
    run.bold = True
    d.add_paragraph('Skills: Python, Django, AWS')
    d.save(buf)
    return buf.getvalue()


def _make_candidate(email, resume_bytes, filename, content_type, raw_text=''):
    user = User.objects.create_user(email=email, password='Password123!', role=User.Role.CANDIDATE)
    return CandidateProfile.objects.create(
        user=user,
        full_name='Jane Smith',
        resume=SimpleUploadedFile(filename, resume_bytes, content_type=content_type),
        original_filename=filename,
        mime_type=content_type,
        raw_resume_text=raw_text,
    )


@pytest.mark.django_db
def test_docx_resume_preview_renders_html():
    profile = _make_candidate(
        'docx_preview@talentvault.in', _docx_bytes(), 'resume.docx', DOCX_CT
    )
    client = Client()
    client.force_login(profile.user)

    response = client.get(reverse('frontend:candidate_resume_preview', kwargs={'pk': profile.pk}))

    assert response.status_code == 200
    assert 'text/html' in response['Content-Type']
    body = response.content.decode('utf-8')
    assert 'Jane Smith' in body
    assert 'Python and AWS' in body
    # Mammoth preserves basic formatting (headings / bold).
    assert '<h1>' in body or '<strong>' in body


@pytest.mark.django_db
def test_corrupted_docx_preview_falls_back_gracefully():
    profile = _make_candidate(
        'docx_corrupt@talentvault.in',
        b'this is definitely not a real docx file',
        'corrupt.docx',
        DOCX_CT,
        raw_text='Extracted fallback text',
    )
    client = Client()
    client.force_login(profile.user)

    response = client.get(reverse('frontend:candidate_resume_preview', kwargs={'pk': profile.pk}))

    # Graceful fallback: never a 500, still 200 HTML with a download option.
    assert response.status_code == 200
    body = response.content.decode('utf-8')
    assert 'Download' in body


@pytest.mark.django_db
def test_public_profile_embeds_docx_iframe_not_unavailable_message():
    profile = _make_candidate(
        'docx_public@talentvault.in', _docx_bytes(), 'resume.docx', DOCX_CT
    )

    response = Client().get(reverse('frontend:public_candidate_profile', kwargs={'pk': profile.pk}))

    assert response.status_code == 200
    body = response.content.decode('utf-8')
    # The DOCX branch now embeds the server-rendered preview.
    assert f'/share/candidate/{profile.pk}/resume-preview/' in body
    assert 'Preview not available for this file type (DOCX)' not in body


@pytest.mark.django_db
def test_pdf_preview_still_works():
    profile = _make_candidate(
        'pdf_preview@talentvault.in', b'%PDF-1.4 test pdf content', 'resume.pdf', PDF_CT
    )
    client = Client()
    client.force_login(profile.user)

    response = client.get(reverse('frontend:candidate_resume_preview', kwargs={'pk': profile.pk}))

    assert response.status_code == 200
    assert 'pdf' in response['Content-Type']


@pytest.mark.django_db
def test_admin_candidate_detail_embeds_docx_iframe_not_unavailable_message():
    admin = User.objects.create_superuser(
        email='docx_admin@talentvault.in', password='Password123!', role=User.Role.SUPER_ADMIN
    )
    cand_user = User.objects.create_user(email='docx_admin_cand@example.com', role=User.Role.CANDIDATE)
    profile = CandidateProfile.objects.create(
        user=cand_user,
        full_name='Jane Smith',
        resume=SimpleUploadedFile('resume.docx', _docx_bytes(), content_type=DOCX_CT),
        original_filename='resume.docx',
        mime_type=DOCX_CT,
        uploaded_by=admin,
        created_by=admin,
    )

    client = Client()
    client.force_login(admin)
    response = client.get(reverse('frontend:candidate_detail', kwargs={'pk': profile.pk}))

    assert response.status_code == 200
    body = response.content.decode('utf-8')
    assert reverse('frontend:candidate_resume_preview', kwargs={'pk': profile.pk}) in body
    assert 'Preview not available for this file type (DOCX)' not in body


def test_admin_preview_template_has_docx_branch():
    src = pathlib.Path('templates/candidate_detail.html').read_text(encoding='utf-8')
    assert "resume_extension == 'docx'" in src
    assert "candidate_resume_preview" in src
